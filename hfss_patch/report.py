"""Results -> metrics.json + figures + Word (.docx) report.

CLI:
  python -m hfss_patch.report --config <yaml> --results <dir> [--out report.docx]
  python -m hfss_patch.report --config <yaml> --results <dir> --synthetic   # preview with fake data
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import plots
from .capture import CAPTURE_DIR, CAPTURES, MANIFEST, load_manifest
from . import postprocess as pp
from .analytic import design_patch, resonant_frequency
from .config import ConfigError, JobConfig, load_config

logger = logging.getLogger(__name__)

FONT = "Malgun Gothic"
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
GOOD = RGBColor(0x0C, 0x7A, 0x0C)
BAD = RGBColor(0xC0, 0x30, 0x30)
HEADER_FILL = "E8EEF6"


@dataclass(frozen=True)
class Spec:
    """Acceptance criteria approved in docs/plan.md (Q3)."""

    s11_max_db: float = -10.0
    band_ghz: tuple[float, float] = pp.ISM_BAND_GHZ
    gain_min_dbi: float = 4.0


@dataclass
class ReportData:
    cfg: JobConfig
    results_dir: Path
    synthetic: bool
    run: dict[str, Any] = field(default_factory=dict)
    s11: pp.S11Metrics | None = None
    pattern: pp.PatternMetrics | None = None
    peak3d: tuple[float, float, float] | None = None
    antenna: dict[str, float] = field(default_factory=dict)
    parametric: dict[str, list[dict[str, float]]] = field(default_factory=dict)
    tuning: dict[str, float | None] = field(default_factory=dict)
    figures: dict[str, Path] = field(default_factory=dict)
    captures: dict[str, Path] = field(default_factory=dict)  # AEDT screen captures
    s11_curve: tuple | None = None  # (freq_ghz, s11_db) for comparison overlays
    capture_status: str = "not run"
    convergence_tail: list[str] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)

    def metrics_dict(self) -> dict[str, Any]:
        return {
            "synthetic": self.synthetic,
            "s11": self.s11.as_dict() if self.s11 else None,
            "pattern": vars(self.pattern) if self.pattern else None,
            "peak3d_dbi_theta_phi": self.peak3d,
            "antenna_params": self.antenna,
            "parametric": self.parametric,
            "tuning": self.tuning,
            "captures": {k: str(v) for k, v in self.captures.items()},
            "capture_status": self.capture_status,
            "missing": self.missing,
        }


# --- analysis ----------------------------------------------------------------------------


def analyze(cfg: JobConfig, results_dir: Path) -> ReportData:
    """Each result file is optional: a missing/broken file becomes a 'missing' entry, not a crash."""
    figdir = results_dir / "figures"
    figdir.mkdir(parents=True, exist_ok=True)
    run: dict[str, Any] = {}
    try:
        run = json.loads((results_dir / "run_summary.json").read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        logger.warning("run_summary.json unavailable: %s", exc)
    rd = ReportData(cfg, results_dir, synthetic=run.get("data_origin") == "SYNTHETIC", run=run)
    if run.get("config_sha256") and run["config_sha256"] != cfg.sha256:
        logger.warning("config sha256 differs from the one used for the run — report may not match results")
        rd.missing.append("config sha256 mismatch (results produced with a different YAML)")

    rd.figures["geometry"] = plots.geometry_figure(cfg, figdir / "geometry.png")

    def attempt(label: str, fn) -> None:
        try:
            fn()
        except (pp.ResultsError, KeyError, ValueError, IndexError) as exc:
            logger.warning("%s: %s", label, exc)
            rd.missing.append(f"{label}: {exc}")

    def do_s11() -> None:
        m, c = pp.s11_from_table(pp.read_table(results_dir / "s11.csv"))
        rd.s11 = m
        rd.s11_curve = (c["freq_ghz"], c["s11_db"])
        rd.figures["s11"] = plots.s11_figure(c["freq_ghz"], c["s11_db"], m, figdir / "s11.png")
        if "z" in c:
            rd.figures["smith"] = plots.smith_figure(c["freq_ghz"], c["z"], m.f_res_ghz, figdir / "smith.png")

    def do_cuts() -> None:
        rd.pattern, cuts = pp.pattern_metrics(pp.read_table(results_dir / "farfield_cuts.csv"))
        rd.figures["pattern"] = plots.pattern_figure(cuts, figdir / "pattern.png")

    def do_3d() -> None:
        t = pp.read_table(results_dir / "farfield_3d.csv")
        rd.peak3d = pp.peak_from_3d(t)
        rd.figures["gain3d"] = plots.gain3d_figure(t.col("Phi"), t.col("Theta"), t.col("dB(GainTotal)"),
                                                   figdir / "gain3d.png")

    def do_ant() -> None:
        rd.antenna = pp.antenna_params(pp.read_table(results_dir / "antenna_params.csv"))

    attempt("s11.csv", do_s11)
    attempt("farfield_cuts.csv", do_cuts)
    attempt("farfield_3d.csv", do_3d)
    attempt("antenna_params.csv", do_ant)

    for pc in cfg.parametrics:
        def do_param(var: str = pc.variable) -> None:
            t = pp.read_table(results_dir / f"s11_param_{var}.csv")
            rows = pp.parametric_summary(t, var)
            rd.parametric[var] = rows
            rd.figures[f"param_{var}"] = plots.parametric_figure(
                var, pp.parametric_curves(t, var), getattr(cfg.patch, var), figdir / f"param_{var}.png"
            )
            if var == "L":
                rd.tuning["L_for_f0_mm"] = pp.tune_parameter(rows, "L", cfg.setup.f0)

        attempt(f"s11_param_{pc.variable}.csv", do_param)

    rd.captures = load_manifest(results_dir)
    try:
        raw = json.loads((results_dir / CAPTURE_DIR / MANIFEST).read_text(encoding="utf-8"))
        bad = [k for k, v in raw.items() if str(v).startswith("FAILED")]
        rd.capture_status = f"{len(rd.captures)}/{len(CAPTURES)} captured" + (f", failed: {', '.join(bad)}" if bad else "")
        rd.missing += [f"capture {k}: {raw[k]}" for k in bad]
    except (OSError, ValueError):
        rd.capture_status = "not run (matplotlib figures only)"

    rd.convergence_tail = pp.read_text_tail(results_dir / "convergence.prop", 15)
    if not rd.convergence_tail:
        rd.missing.append("convergence.prop")
    return rd


def gain_dbi(rd: ReportData) -> float | None:
    if "PeakGain_dBi" in rd.antenna and math.isfinite(rd.antenna["PeakGain_dBi"]):
        return rd.antenna["PeakGain_dBi"]
    if rd.peak3d:
        return rd.peak3d[0]
    return None


def verdicts(rd: ReportData, spec: Spec) -> list[tuple[str, str, str, bool | None]]:
    """(item, target, result, pass?) — pass None = informative only."""
    rows: list[tuple[str, str, str, bool | None]] = []
    lo, hi = spec.band_ghz
    if rd.s11:
        m = rd.s11
        rows.append(("공진 주파수", f"{lo:.4g}–{hi:.4g} GHz 내", f"{m.f_res_ghz:.3f} GHz", lo <= m.f_res_ghz <= hi))
        rows.append(("최소 |S11|", f"≤ {spec.s11_max_db:.0f} dB", f"{m.s11_min_db:.1f} dB", m.s11_min_db <= spec.s11_max_db))
        bw = f"{m.bw10_mhz:.1f} MHz ({m.bw10_pct:.2f} %)" if m.bw10_mhz else "미정의 (-10 dB 미도달/범위 밖)"
        rows.append(("-10 dB 대역폭", f"≥ {(hi - lo) * 1e3:.1f} MHz (ISM)", bw,
                     bool(m.bw10_mhz and m.bw10_mhz >= (hi - lo) * 1e3)))
        rows.append(("ISM 대역 내 최악 |S11|", f"≤ {spec.s11_max_db:.0f} dB", f"{m.ism_worst_s11_db:.1f} dB", m.ism_pass))
    g = gain_dbi(rd)
    if g is not None:
        rows.append(("Peak Gain", f"≥ {spec.gain_min_dbi:.1f} dBi", f"{g:.2f} dBi", g >= spec.gain_min_dbi))
    if "RadiationEfficiency_pct" in rd.antenna:
        rows.append(("방사 효율", "참고", f"{rd.antenna['RadiationEfficiency_pct']:.1f} %", None))
    if rd.pattern and rd.pattern.xpd_boresight_db is not None:
        rows.append(("교차편파 분리도 (boresight)", "참고", f"{rd.pattern.xpd_boresight_db:.1f} dB", None))
    return rows


def recommendations(rd: ReportData, spec: Spec) -> list[str]:
    out: list[str] = []
    cfg = rd.cfg
    if rd.s11:
        df = rd.s11.f_res_ghz - cfg.setup.f0
        if abs(df) > 0.005:
            l_new = rd.tuning.get("L_for_f0_mm")
            hint = f" 파라메트릭 결과 기준 L ≈ {l_new:.2f} mm로 조정하면 {cfg.setup.f0} GHz에 맞출 수 있습니다." if l_new else ""
            out.append(f"공진 주파수가 목표 대비 {df * 1e3:+.0f} MHz 벗어나 있습니다.{hint}")
        if not rd.s11.ism_pass:
            out.append(
                f"ISM 전대역(83.5 MHz) |S11| ≤ {spec.s11_max_db:.0f} dB를 만족하지 못합니다. 단일 패치의 대역폭은 "
                f"기판 두께/유전율로 제한되며(현재 εr {cfg.substrate.er}, h {cfg.substrate.h} mm), 같은 두께에서 "
                "저손실 기판으로 바꾸면 손실에 의한 대역 확장이 사라져 오히려 대역폭이 줄어듭니다. 전대역 확보에는 "
                "공기/폼 기판(h ≈ 5 mm급), 적층(stacked) 패치, 또는 U-slot 구조 변경이 필요합니다."
            )
    g = gain_dbi(rd)
    if g is not None and g < spec.gain_min_dbi:
        eff = rd.antenna.get("RadiationEfficiency_pct")
        e = f"(방사 효율 {eff:.0f} %)" if eff else ""
        out.append(
            f"Peak Gain {g:.2f} dBi로 목표 {spec.gain_min_dbi:.1f} dBi에 미달합니다{e}. FR-4의 유전 손실(tanδ "
            f"{cfg.substrate.tand})이 주 원인이며 저손실 기판(예: RO4003C, tanδ 0.0027) 적용 시 효율·이득 개선이 "
            "예상됩니다(대역폭은 감소)."
        )
    if rd.s11 and rd.s11.s11_min_db > spec.s11_max_db:
        out.append("임피던스 정합이 부족합니다. inset 깊이 y0 파라메트릭 결과를 참고해 y0를 재조정하십시오.")
    if not out:
        out.append("모든 승인 기준을 만족합니다. 제작 공차(±0.1 mm 식각, εr ±0.2) 민감도 확인 후 시제작을 권고합니다.")
    return out


def substrate_label(cfg: JobConfig) -> str:
    s = cfg.substrate
    return f"{s.material.removeprefix('SUB_')} (εr {s.er:g}, h {s.h:g} mm)"


def comparison_rows(rds: list[ReportData], spec: Spec) -> list[list[Any]]:
    """Metric rows for side-by-side substrate comparison (one column per design)."""

    def cell(ok: bool | None, text: str):
        return (text, GOOD if ok else BAD) if ok is not None else text

    rows: list[list[Any]] = [
        ["기판"] + [substrate_label(r.cfg) for r in rds],
        ["패치 W × L"] + [f"{r.cfg.patch.W:.2f} × {r.cfg.patch.L:.2f} mm" for r in rds],
        ["Inset y0 / Wf"] + [f"{r.cfg.patch.y0:.2f} / {r.cfg.patch.Wf:.2f} mm" for r in rds],
    ]
    lo, hi = spec.band_ghz

    def s11_row(label: str, fn) -> None:
        rows.append([label] + [fn(r.s11) if r.s11 else "—" for r in rds])

    s11_row("공진 주파수", lambda m: cell(lo <= m.f_res_ghz <= hi, f"{m.f_res_ghz:.3f} GHz"))
    s11_row("최소 |S11|", lambda m: cell(m.s11_min_db <= spec.s11_max_db, f"{m.s11_min_db:.1f} dB"))
    s11_row("-10 dB 대역폭", lambda m: cell(bool(m.bw10_mhz and m.bw10_mhz >= (hi - lo) * 1e3),
                                            _fmt(m.bw10_mhz, 1, " MHz")))
    s11_row("ISM 최악 |S11|", lambda m: cell(m.ism_pass, f"{m.ism_worst_s11_db:.1f} dB"))
    gains = [gain_dbi(r) for r in rds]
    rows.append(["Peak Gain"] + [cell(g >= spec.gain_min_dbi, f"{g:.2f} dBi") if g is not None else "—" for g in gains])
    rows.append(["방사 효율"] + [_fmt(r.antenna.get("RadiationEfficiency_pct"), 1, " %") for r in rds])
    rows.append(["데이터"] + ["SYNTHETIC" if r.synthetic else "HFSS" for r in rds])
    return rows


def comparison_findings(base: ReportData, alt: ReportData, spec: Spec) -> list[str]:
    out: list[str] = []
    gb, ga = gain_dbi(base), gain_dbi(alt)
    if gb is not None and ga is not None:
        out.append(f"이득: {substrate_label(base.cfg)} {gb:.2f} dBi → {substrate_label(alt.cfg)} {ga:.2f} dBi "
                   f"({ga - gb:+.2f} dB). 목표 {spec.gain_min_dbi:.0f} dBi "
                   f"{'충족' if ga >= spec.gain_min_dbi else '미달'}.")
    if base.s11 and alt.s11 and base.s11.bw10_mhz and alt.s11.bw10_mhz:
        out.append(f"-10 dB 대역폭: {base.s11.bw10_mhz:.1f} → {alt.s11.bw10_mhz:.1f} MHz. 저손실 기판은 손실에 의한 "
                   "대역 확장이 없어 대역폭이 줄어듭니다.")
    if alt.s11 and not alt.s11.ism_pass:
        out.append("두 안 모두 ISM 전대역 기준을 만족하지 못합니다. 전대역이 필수 요구라면 구조 변경(공기/폼 기판, "
                   "적층 패치, U-slot)이 필요하며, 채널 단위 운용(예: 2.44 GHz 중심 ±15 MHz)이라면 대안 기판으로 충분합니다.")
    return out


# --- docx helpers -------------------------------------------------------------------------


def _set_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _para(doc, text: str = "", size: float = 10, bold: bool = False, color: RGBColor | None = None,
          align=None, style: str | None = None):
    p = doc.add_paragraph(style=style)
    if text:
        _set_font(p.add_run(text), size, bold, color)
    if align is not None:
        p.alignment = align
    return p


def _heading(doc, text: str, level: int) -> None:
    h = doc.add_heading(level=level)
    _set_font(h.add_run(text), color=INK)


def _shade(cell, fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _table(doc, header: list[str], rows: list[list[Any]], widths_cm: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _set_font(c.paragraphs[0].add_run(h), 9, True)
        _shade(c, HEADER_FILL)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            color = None
            if isinstance(v, tuple):  # (text, color)
                v, color = v
            _set_font(cells[i].paragraphs[0].add_run(str(v)), 9, color=color, bold=color is not None)
    if widths_cm:
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


class _Figures:
    """Sequential figure numbering; a missing image does not consume a number."""

    def __init__(self) -> None:
        self.n = 0

    def add(self, doc, path: Path | None, caption: str, width_cm: float = 15.0, optional: bool = False) -> bool:
        if not path or not Path(path).exists():
            if not optional:
                _para(doc, f"[그림 없음: {caption}]", 9, color=INK2)
            return False
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.n += 1
        _para(doc, f"그림 {self.n}. {caption}", 9, color=INK2, align=WD_ALIGN_PARAGRAPH.CENTER)
        return True


def _fmt(v: float | None, nd: int = 2, unit: str = "") -> str:
    return "—" if v is None or (isinstance(v, float) and not math.isfinite(v)) else f"{v:.{nd}f}{unit}"


# --- document ------------------------------------------------------------------------------


def build_docx(rd: ReportData, out: Path, spec: Spec = Spec(), customer: str = "", doc_no: str = "",
               alt: ReportData | None = None) -> Path:
    cfg = rd.cfg
    synthetic = rd.synthetic or (alt is not None and alt.synthetic)  # any fake data -> watermark everything
    fig = _Figures()
    cap = rd.captures
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

    hdr = sec.header.paragraphs[0]
    if synthetic:
        _set_font(hdr.add_run("SYNTHETIC DATA — 파이프라인 미리보기용, HFSS 해석 결과 아님 / NOT FOR DELIVERY"), 9, True, BAD)
    else:
        _set_font(hdr.add_run(f"HFSS 해석 보고서 — {cfg.project.name}"), 8, color=INK2)
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    _para(doc, "2.4 GHz 마이크로스트립 패치 안테나", 22, True, INK, WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "HFSS 설계 및 전자기 해석 보고서", 16, False, INK2, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    if synthetic:
        _para(doc, "※ 본 문서는 합성(synthetic) 데이터로 생성한 양식 미리보기입니다. 수치는 실제 해석 결과가 아닙니다.",
              11, True, BAD, WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(6):
        doc.add_paragraph()
    _table(
        doc,
        ["항목", "내용"],
        [
            ["의뢰처", customer or "—"],
            ["문서 번호", doc_no or "—"],
            ["작성일", date.today().isoformat()],
            ["해석 도구", f"Ansys HFSS {rd.run.get('aedt_version', '—')} (PyAEDT 자동화)"],
            ["데이터 출처", "SYNTHETIC (미리보기)" if synthetic else "HFSS 해석 결과"],
        ],
        [4, 11],
    )
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    # 1. Summary
    _heading(doc, "1. 요약", 1)
    _para(doc, f"요청 사양: HFSS를 이용한 {cfg.setup.f0} GHz 공진 패치 안테나 설계, 해석 및 결과 보고. "
               "아래 표는 승인된 성능 목표 대비 해석 결과의 판정입니다.")
    vrows = []
    for item, tgt, res, ok in verdicts(rd, spec):
        mark = ("적합", GOOD) if ok else ("부적합", BAD) if ok is False else ("참고", None)
        vrows.append([item, tgt, res, mark if mark[1] else mark[0]])
    _table(doc, ["항목", "목표", "결과", "판정"], vrows, [4.5, 4, 4.5, 2.5])
    if alt is not None:
        _para(doc, f"대안 기판 {substrate_label(alt.cfg)} 비교 결과는 8장에 정리하였습니다.", 9, color=INK2)
    _heading(doc, "주요 권고사항", 2)
    for r in recommendations(rd, spec):
        _para(doc, r, style="List Bullet")

    # 2. Assumptions
    _heading(doc, "2. 설계 가정사항", 1)
    _para(doc, "의뢰서에 기판, 급전 방식, 성능 목표가 명시되지 않아 다음과 같이 가정하였습니다.")
    s = cfg.substrate
    _table(
        doc,
        ["항목", "가정"],
        [
            ["기판", f"FR-4 계열, εr {s.er}, tanδ {s.tand}, 두께 {s.h} mm"],
            ["도체", "완전 도체(PerfE) 시트 — 동박 두께·도체 손실 무시"],
            ["급전", "Inset microstrip feed, 50 Ω"],
            ["편파", "선형 편파 (급전 방향 y축, E-plane = φ 90°)"],
            ["GND/기판 크기", f"{cfg.board.Wsub:.0f} × {cfg.board.Lsub:.0f} mm"],
            ["성능 목표", f"ISM {spec.band_ghz[0]}–{spec.band_ghz[1]} GHz |S11| ≤ {spec.s11_max_db:.0f} dB, "
                         f"Gain ≥ {spec.gain_min_dbi:.0f} dBi"],
        ],
        [4, 11],
    )

    # 3. Design
    _heading(doc, "3. 안테나 설계", 1)
    _heading(doc, "3.1 해석식 초기 설계 (전송선로 모델)", 2)
    d = design_patch(cfg.setup.f0, s.er, s.h)
    _table(
        doc,
        ["파라미터", "값", "비고"],
        [
            ["패치 폭 W", f"{d.W_mm:.2f} mm", "c/(2f0)·√(2/(εr+1))"],
            ["유효 유전율 ε_eff", f"{d.eps_eff:.3f}", "Hammerstad"],
            ["Fringing ΔL", f"{d.dL_mm:.3f} mm", ""],
            ["패치 길이 L", f"{d.L_mm:.2f} mm", "c/(2f0√ε_eff) − 2ΔL"],
            ["Edge 입력저항", f"{d.Rin_edge_ohm:.0f} Ω", "G1, G12 적분식 (Balanis)"],
            ["Inset 깊이 y0", f"{d.y0_mm:.2f} mm", "Rin·cos²(πy0/L) = 50 Ω"],
            ["50 Ω 선로 폭 Wf", f"{d.Wf_mm:.2f} mm", "Hammerstad 합성"],
            ["VSWR<2 대역폭 추정", f"{d.bw_vswr2_pct:.2f} %", "무손실 기준 (Jackson)"],
        ],
        [5, 4, 6],
    )
    _heading(doc, "3.2 해석 모델 치수", 2)
    p = cfg.patch
    _table(
        doc,
        ["W", "L", "y0", "gap", "Wf", "h", "기판"],
        [[f"{p.W:.2f}", f"{p.L:.2f}", f"{p.y0:.2f}", f"{p.gap:.2f}", f"{p.Wf:.2f}", f"{s.h:.2f}",
          f"{cfg.board.Wsub:.0f}×{cfg.board.Lsub:.0f}"]],
    )
    fig.add(doc, rd.figures.get("geometry"), "안테나 형상 및 치수 (상면도, 단위 mm)", 11)
    for k in ("model_iso", "model_top"):
        fig.add(doc, cap.get(k), CAPTURES[k][1], 14, optional=True)

    # 4. Setup
    _heading(doc, "4. 해석 조건", 1)
    st = cfg.setup
    _table(
        doc,
        ["항목", "설정"],
        [
            ["Solution type", "HFSS Driven Modal"],
            ["Port", "Wave Port (기판 엣지, 폭 " f"{cfg.port.w_factor:g}·Wf, 높이 {cfg.port.h_factor:g}·h), 50 Ω renormalize"],
            ["경계 조건", f"Radiation box, 이격 {cfg.air:g} mm (≥ λ0/4 @ {st.sweep_start} GHz)"],
            ["Adaptive mesh", f"{st.f0} GHz, Max ΔS {st.max_delta_s}, 최대 {st.max_passes} pass"],
            ["주파수 스윕", f"{st.sweep_type}, {st.sweep_start}–{st.sweep_stop} GHz, step {st.sweep_step * 1e3:g} MHz"],
            ["원거리장", f"Infinite sphere, θ 0–180°/{cfg.farfield.theta_step:g}°, φ 0–360°/{cfg.farfield.phi_step:g}°"],
            ["파라메트릭", ", ".join(f"{pc.variable} {pc.start:g}–{pc.stop:g} mm/{pc.step:g}" for pc in cfg.parametrics) or "—"],
        ],
        [4, 11],
    )
    if rd.convergence_tail:
        _para(doc, "수렴 이력 (convergence.prop 발췌)", 9, True)
        for line in rd.convergence_tail:
            _set_font(doc.add_paragraph().add_run(line), 8, color=INK2)

    # 5. Results
    _heading(doc, "5. 해석 결과", 1)
    _heading(doc, "5.1 반사 계수 및 입력 임피던스", 2)
    if rd.s11:
        m = rd.s11
        z = f"{m.z_res_ohm.real:.1f} {m.z_res_ohm.imag:+.1f}j Ω" if m.z_res_ohm is not None else "—"
        _table(
            doc,
            ["공진 주파수", "최소 |S11|", "-10 dB 대역", "대역폭", "Z @ 공진", "VSWR @ 공진"],
            [[f"{m.f_res_ghz:.3f} GHz", f"{m.s11_min_db:.1f} dB",
              f"{_fmt(m.bw10_lo_ghz, 3)}–{_fmt(m.bw10_hi_ghz, 3)} GHz", _fmt(m.bw10_mhz, 1, " MHz"), z,
              _fmt(m.vswr_res, 2)]],
        )
    fig.add(doc, rd.figures.get("s11"), "반사 계수 |S11| (음영: ISM 대역)")
    fig.add(doc, rd.figures.get("smith"), "입력 임피던스 Smith chart", 10)

    _heading(doc, "5.2 방사 특성", 2)
    if rd.pattern:
        pm = rd.pattern
        _table(
            doc,
            ["Boresight Gain", "Peak Gain (3D)", "HPBW E/H", "F/B", "XPD", "효율"],
            [[f"{pm.boresight_gain_dbi:.2f} dBi", _fmt(rd.peak3d[0] if rd.peak3d else None, 2, " dBi"),
              f"{_fmt(pm.hpbw_e_deg, 0, '°')} / {_fmt(pm.hpbw_h_deg, 0, '°')}", f"{pm.front_to_back_db:.1f} dB",
              _fmt(pm.xpd_boresight_db, 1, " dB"), _fmt(rd.antenna.get("RadiationEfficiency_pct"), 1, " %")]],
        )
    if rd.antenna:
        _para(doc, "Antenna Parameters: " + ", ".join(f"{k} = {v:.2f}" for k, v in rd.antenna.items()), 9, color=INK2)
    fig.add(doc, rd.figures.get("pattern"), "2D 방사 패턴 (E-plane φ=90°, H-plane φ=0°)")
    fig.add(doc, cap.get("pattern3d"), CAPTURES["pattern3d"][1], 12, optional=True)
    fig.add(doc, rd.figures.get("gain3d"), "3D 이득 분포 (θ–φ 맵)")

    _heading(doc, "5.3 전류 및 전계 분포", 2)
    if not any(k in cap for k in ("jsurf", "efield_cut")):
        _para(doc, "AEDT 필드 캡처가 없습니다 (캡처 단계 미실행 또는 실패 — 부록의 캡처 상태 참조).", 9, color=INK2)
    else:
        _para(doc, "TM010 모드에서는 표면 전류가 급전 방향(y)으로 흐르고, 전계는 두 방사 엣지(y = ±L/2)에서 최대가 됩니다.",
              9, color=INK2)
    for k in ("jsurf", "efield_cut"):
        fig.add(doc, cap.get(k), CAPTURES[k][1], 14, optional=True)

    _heading(doc, "5.4 파라메트릭 민감도", 2)
    for var, rows in rd.parametric.items():
        _table(
            doc,
            [f"{var} [mm]", "공진 주파수 [GHz]", "최소 |S11| [dB]", "-10 dB BW [MHz]"],
            [[f"{r[var]:g}", f"{r['f_res_ghz']:.3f}", f"{r['s11_min_db']:.1f}", _fmt(r["bw10_mhz"], 1)] for r in rows],
        )
        fig.add(doc, rd.figures.get(f"param_{var}"), f"{var} 파라메트릭 스윕 |S11|")
    if rd.tuning.get("L_for_f0_mm"):
        _para(doc, f"L 민감도 선형 근사 결과, f0 = {cfg.setup.f0} GHz 공진을 위한 L ≈ {rd.tuning['L_for_f0_mm']:.2f} mm.")

    # 6. Analytic cross-check
    _heading(doc, "6. 해석식 대비 검증", 1)
    if rd.s11:
        f_tl = resonant_frequency(p.L, p.W, s.er, s.h)
        err = (rd.s11.f_res_ghz - f_tl) / f_tl * 100
        _table(
            doc,
            ["항목", "전송선로 모델", "HFSS", "차이"],
            [["공진 주파수", f"{f_tl:.3f} GHz", f"{rd.s11.f_res_ghz:.3f} GHz", f"{err:+.2f} %"]],
        )
        _para(doc, "전송선로 모델은 inset 노치·유한 GND·급전선 영향을 반영하지 않으므로 수 % 이내의 차이는 정상 범위입니다.",
              9, color=INK2)

    # 7. Conclusion
    _heading(doc, "7. 결론", 1)
    for r in recommendations(rd, spec):
        _para(doc, r, style="List Bullet")

    # Appendix
    if alt is not None:
        _heading(doc, f"8. 대안 비교: {substrate_label(alt.cfg)}", 1)
        _para(doc, f"원안({substrate_label(cfg)})과 동일한 해석 조건으로 대안 기판을 설계·해석하여 비교하였습니다. "
                   "치수는 각 기판에 대해 전송선로 모델로 재설계한 값입니다.")
        _table(doc, ["항목", "원안", "대안"], comparison_rows([rd, alt], spec), [4, 5.5, 5.5])
        if rd.s11_curve and alt.s11_curve:
            path = rd.results_dir / "figures" / "compare_s11.png"
            plots.compare_s11_figure(
                [(substrate_label(rd.cfg), *rd.s11_curve), (substrate_label(alt.cfg), *alt.s11_curve)], path
            )
            fig.add(doc, path, "원안/대안 |S11| 비교")
        for line in comparison_findings(rd, alt, spec):
            _para(doc, line, style="List Bullet")
        if alt.missing:
            _para(doc, "대안 결과 누락/경고: " + "; ".join(alt.missing), 9, color=INK2)

    _heading(doc, "부록. 실행 정보 (재현성)", 1)
    _table(
        doc,
        ["항목", "값"],
        [
            ["설정 파일", Path(cfg.source_path).name],
            ["설정 SHA-256", cfg.sha256[:16] + "…"],
            ["AEDT 버전", str(rd.run.get("aedt_version", "—"))],
            ["실행 상태", str(rd.run.get("status", "—"))],
            ["AEDT 캡처", rd.capture_status],
            ["누락/경고", "; ".join(rd.missing) or "없음"],
        ],
        [4, 11],
    )
    if "s11_report" in cap:
        _para(doc, "AEDT 원본 리포트 화면 (수치 교차 확인용)", 9, True)
        fig.add(doc, cap["s11_report"], CAPTURES["s11_report"][1], 14)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build the HFSS patch antenna Word report")
    ap.add_argument("--config", required=True, type=Path)
    ap.add_argument("--results", type=Path, help="results dir (default: project.output_dir)")
    ap.add_argument("--out", type=Path, help="output .docx (default: <results>/report.docx)")
    ap.add_argument("--synthetic", action="store_true", help="first write SYNTHETIC data into --results (preview)")
    ap.add_argument("--compare", type=Path, help="alternative design config to compare against (e.g. RO4003C)")
    ap.add_argument("--compare-results", type=Path, help="results dir of the comparison (default: its output_dir)")
    ap.add_argument("--customer", default="")
    ap.add_argument("--doc-no", default="")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-7s %(name)s: %(message)s")

    try:
        cfg = load_config(args.config)
    except ConfigError as exc:
        print(f"CONFIG ERROR: {exc}", file=sys.stderr)
        return 2
    results = (args.results or Path(cfg.project.output_dir)).resolve()
    alt_cfg = None
    if args.compare:
        try:
            alt_cfg = load_config(args.compare)
        except ConfigError as exc:
            print(f"CONFIG ERROR (--compare): {exc}", file=sys.stderr)
            return 2
    alt_results = (args.compare_results or Path(alt_cfg.project.output_dir)).resolve() if alt_cfg else None
    if alt_results is not None and alt_results == results:
        logger.error("--compare results dir must differ from the baseline results dir")
        return 2

    if args.synthetic:
        from .synthetic import write_synthetic_results

        for c, r in ((cfg, results), (alt_cfg, alt_results)):
            if c is not None:
                write_synthetic_results(c, r)
                logger.warning("SYNTHETIC data written to %s — preview only", r)
    for r in (results, alt_results):
        if r is not None and not r.is_dir():
            logger.error("results dir not found: %s", r)
            return 2

    rd = analyze(cfg, results)
    alt = analyze(alt_cfg, alt_results) if alt_cfg else None
    (results / "metrics.json").write_text(json.dumps(rd.metrics_dict(), indent=2, default=str), encoding="utf-8")
    if alt is not None:
        (alt_results / "metrics.json").write_text(json.dumps(alt.metrics_dict(), indent=2, default=str),
                                                  encoding="utf-8")
    out = build_docx(rd, args.out or results / "report.docx", customer=args.customer, doc_no=args.doc_no, alt=alt)
    missing = rd.missing + ([f"[compare] {m}" for m in alt.missing] if alt else [])
    logger.info("report written: %s (missing: %s)", out, missing or "none")
    return 0 if not missing else 4


if __name__ == "__main__":
    raise SystemExit(main())
