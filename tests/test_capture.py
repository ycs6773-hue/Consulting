"""Capture stage against a fake PyAEDT: isolation, session release, CLI exit codes, report insertion."""

from __future__ import annotations

import json
import sys
import types
from dataclasses import replace
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from docx import Document

from hfss_patch import run as run_cli
from hfss_patch.builder import HfssJobError
from hfss_patch.capture import CAPTURES, load_manifest, run_capture
from hfss_patch.config import ConfigError, load_config
from hfss_patch.report import analyze, build_docx
from hfss_patch.synthetic import write_synthetic_results

CFG = Path(__file__).resolve().parents[1] / "hfss_patch" / "configs" / "patch_2g4_fr4.yaml"


def _touch(path, *a, **k):
    Path(path).write_bytes(b"\xff\xd8fakejpg")
    return str(path)


def make_hfss(fail: set[str] = frozenset()) -> MagicMock:
    """Fake Hfss whose image exports really write files (unless the capture is in `fail`)."""
    h = MagicMock(name="Hfss")
    h.modeler.planes = []

    def model_pic(full_name, orientation, **k):
        if f"model_{'iso' if orientation == 'isometric' else 'top'}" in fail:
            return False
        return _touch(full_name)

    h.post.export_model_picture.side_effect = model_pic

    def fieldplot(name):
        plot = MagicMock()
        plot.export_image.side_effect = (lambda *a, **k: False) if name in fail else _touch
        return plot

    h.post.create_fieldplot_surface.side_effect = lambda *a, **k: fieldplot("jsurf")
    h.post.create_fieldplot_cutplane.side_effect = lambda *a, **k: fieldplot("efield_cut")

    def report_jpg(path, plot_name, **k):
        key = "s11_report" if plot_name == "Cap_S11" else "pattern3d"
        if key in fail:
            raise RuntimeError("report export boom")
        return _touch(path)

    h.post.export_report_to_jpg.side_effect = report_jpg
    return h


@pytest.fixture()
def cfg(tmp_path):
    c = load_config(CFG)
    out = tmp_path / "out"
    out.mkdir()
    (out / f"{c.project.name}.aedt").write_text("solved project stub")
    return replace(c, project=replace(c.project, output_dir=str(out)))


@pytest.fixture()
def fake_aedt(monkeypatch):
    state: dict = {"fail": set(), "instances": [], "kwargs": []}

    def factory(**kwargs):
        if state.get("open_error"):
            raise RuntimeError("cannot start AEDT")
        h = make_hfss(state["fail"])
        state["instances"].append(h)
        state["kwargs"].append(kwargs)
        return h

    mod = types.ModuleType("ansys.aedt.core")
    mod.Hfss = factory
    for name in ("ansys", "ansys.aedt"):
        monkeypatch.setitem(sys.modules, name, types.ModuleType(name))
    monkeypatch.setitem(sys.modules, "ansys.aedt.core", mod)
    return state


def test_capture_all_ok(fake_aedt, cfg):
    manifest = run_capture(cfg)
    assert set(manifest) == set(CAPTURES)
    assert all(Path(v).exists() for v in manifest.values())
    kw = fake_aedt["kwargs"][0]
    assert kw["non_graphical"] is False and kw["new_desktop"] is True and kw["remove_lock"] is True
    h = fake_aedt["instances"][0]
    h.release_desktop.assert_called_once_with(close_projects=True, close_desktop=True)
    h.modeler.create_plane.assert_called_once()
    assert len(load_manifest(Path(cfg.project.output_dir))) == len(CAPTURES)


def test_capture_failures_isolated(fake_aedt, cfg):
    fake_aedt["fail"] = {"jsurf", "pattern3d"}
    manifest = run_capture(cfg)
    assert manifest["jsurf"].startswith("FAILED") and manifest["pattern3d"].startswith("FAILED")
    assert not manifest["model_iso"].startswith("FAILED") and not manifest["s11_report"].startswith("FAILED")
    fake_aedt["instances"][0].release_desktop.assert_called_once()
    assert set(load_manifest(Path(cfg.project.output_dir))) == set(CAPTURES) - {"jsurf", "pattern3d"}


def test_capture_missing_project(fake_aedt, cfg):
    (Path(cfg.project.output_dir) / f"{cfg.project.name}.aedt").unlink()
    with pytest.raises(HfssJobError, match="project not found"):
        run_capture(cfg)
    assert fake_aedt["instances"] == []


def test_capture_open_failure_raises(fake_aedt, cfg):
    fake_aedt["open_error"] = True
    with pytest.raises(HfssJobError, match="open graphical"):
        run_capture(cfg)
    manifest = json.loads((Path(cfg.project.output_dir) / "captures" / "captures.json").read_text())
    assert manifest == {}


def _cli_cfg(tmp_path, cfg) -> Path:
    import yaml

    raw = yaml.safe_load(CFG.read_text(encoding="utf-8"))
    raw["project"]["output_dir"] = cfg.project.output_dir
    p = tmp_path / "c.yaml"
    p.write_text(yaml.safe_dump(raw), encoding="utf-8")
    return p


def test_cli_capture_only_partial_exit4(fake_aedt, cfg, tmp_path):
    fake_aedt["fail"] = {"efield_cut"}
    assert run_cli.main(["--config", str(_cli_cfg(tmp_path, cfg)), "--capture-only"]) == 4


def test_cli_capture_open_failure_after_solve_is_partial(fake_aedt, cfg, tmp_path, monkeypatch):
    monkeypatch.setattr("hfss_patch.builder.run_job", lambda *a, **k: {"status": "ok", "exports": {}})

    def boom(_cfg):
        raise HfssJobError("stage 'capture: open graphical' failed: no display")

    monkeypatch.setattr("hfss_patch.capture.run_capture", boom)
    assert run_cli.main(["--config", str(_cli_cfg(tmp_path, cfg))]) == 4  # solve OK -> not exit 1


def test_cli_no_capture_skips(fake_aedt, cfg, tmp_path, monkeypatch):
    monkeypatch.setattr("hfss_patch.builder.run_job", lambda *a, **k: {"status": "ok", "exports": {}})
    called = []
    monkeypatch.setattr("hfss_patch.capture.run_capture", lambda c: called.append(1) or {})
    assert run_cli.main(["--config", str(_cli_cfg(tmp_path, cfg)), "--no-capture"]) == 0
    assert called == []


def test_capture_config_defaults_and_validation(tmp_path):
    import yaml

    raw = yaml.safe_load(CFG.read_text(encoding="utf-8"))
    raw.pop("capture")
    p = tmp_path / "a.yaml"
    p.write_text(yaml.safe_dump(raw), encoding="utf-8")
    assert load_config(p).capture.enabled is True
    raw["capture"] = {"width": 50}
    p.write_text(yaml.safe_dump(raw), encoding="utf-8")
    with pytest.raises(ConfigError, match="capture"):
        load_config(p)


def test_report_includes_captures_in_order(tmp_path):
    cfg = load_config(CFG)
    out = write_synthetic_results(cfg, tmp_path / "r")
    rd = analyze(cfg, out)
    assert set(rd.captures) == set(CAPTURES) and rd.capture_status.startswith("6/6")
    doc = Document(str(build_docx(rd, tmp_path / "r.docx")))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("그림 ")]
    assert [int(c.split(".")[0].split()[1]) for c in captions] == list(range(1, len(captions) + 1))
    assert len(doc.inline_shapes) == 8 + len(CAPTURES)
    assert any("표면 전류" in c for c in captions)


def test_report_without_captures(tmp_path):
    cfg = load_config(CFG)
    out = write_synthetic_results(cfg, tmp_path / "r", captures=False)
    rd = analyze(cfg, out)
    assert rd.captures == {} and rd.capture_status.startswith("not run") and rd.missing == []
    doc = Document(str(build_docx(rd, tmp_path / "r.docx")))
    assert len(doc.inline_shapes) == 8
    assert "AEDT 필드 캡처가 없습니다" in "\n".join(p.text for p in doc.paragraphs)
