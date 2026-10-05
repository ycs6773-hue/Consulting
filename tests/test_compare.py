"""RO4003C alternative: config consistency, physics expectations, comparison report."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from hfss_patch.analytic import design_patch
from hfss_patch.config import load_config
from hfss_patch.geometry import build_model, check_model
from hfss_patch.report import Spec, analyze, build_docx, comparison_rows, gain_dbi, main, recommendations
from hfss_patch.synthetic import write_synthetic_results

CFGDIR = Path(__file__).resolve().parents[1] / "hfss_patch" / "configs"
FR4 = CFGDIR / "patch_2g4_fr4.yaml"
RO = CFGDIR / "patch_2g4_ro4003c.yaml"


@pytest.mark.parametrize("path", [FR4, RO])
def test_config_matches_analytic_and_geometry(path):
    cfg = load_config(path)
    d = design_patch(cfg.setup.f0, cfg.substrate.er, cfg.substrate.h)
    assert cfg.patch.W == pytest.approx(d.W_mm, abs=0.1)
    assert cfg.patch.L == pytest.approx(d.L_mm, abs=0.1)
    assert cfg.patch.y0 == pytest.approx(d.y0_mm, abs=0.1)
    assert cfg.patch.Wf == pytest.approx(d.Wf_mm, abs=0.05)
    for pc in cfg.parametrics:
        nominal = getattr(cfg.patch, pc.variable)
        assert pc.start <= nominal <= pc.stop, pc.name
    assert check_model(build_model(cfg)) == []


def test_configs_do_not_collide():
    a, b = load_config(FR4), load_config(RO)
    assert a.project.output_dir != b.project.output_dir and a.project.name != b.project.name


@pytest.fixture(scope="module")
def pair(tmp_path_factory):
    base = tmp_path_factory.mktemp("cmp")
    a, b = load_config(FR4), load_config(RO)
    return analyze(a, write_synthetic_results(a, base / "fr4")), analyze(b, write_synthetic_results(b, base / "ro"))


def test_low_loss_tradeoff(pair):
    """Lower tand: higher gain/efficiency but narrower BW (loss no longer broadens it)."""
    fr4, ro = pair
    assert gain_dbi(ro) > gain_dbi(fr4) + 2
    assert gain_dbi(ro) >= Spec().gain_min_dbi > gain_dbi(fr4)
    assert ro.s11.bw10_mhz < fr4.s11.bw10_mhz
    assert not ro.s11.ism_pass and not fr4.s11.ism_pass


def test_recommendation_text_is_physically_correct(pair):
    fr4, _ = pair
    text = " ".join(recommendations(fr4, Spec()))
    assert "대역폭이 줄어듭니다" in text and "효율·이득 개선" in text
    assert "두꺼운 저유전율 기판(예: RO4003C" not in text  # old, wrong bandwidth advice


def test_comparison_rows_shape(pair):
    rows = comparison_rows(list(pair), Spec())
    assert all(len(r) == 3 for r in rows)
    assert rows[0][1].startswith("FR4") and rows[0][2].startswith("RO4003C")


def test_comparison_report(pair, tmp_path):
    fr4, ro = pair
    doc = Document(str(build_docx(fr4, tmp_path / "c.docx", alt=ro)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "8. 대안 비교" in text and "원안/대안 |S11| 비교" in text
    assert "NOT FOR DELIVERY" in doc.sections[0].header.paragraphs[0].text


def test_watermark_if_only_alt_is_synthetic(pair, tmp_path):
    fr4, ro = pair
    fr4.synthetic = False
    try:
        doc = Document(str(build_docx(fr4, tmp_path / "w.docx", alt=ro)))
        assert "SYNTHETIC" in doc.sections[0].header.paragraphs[0].text
    finally:
        fr4.synthetic = True


def test_cli_compare_synthetic(tmp_path):
    rc = main(["--config", str(FR4), "--results", str(tmp_path / "a"), "--compare", str(RO),
               "--compare-results", str(tmp_path / "b"), "--synthetic"])
    assert rc == 0
    assert (tmp_path / "a" / "report.docx").exists() and (tmp_path / "b" / "metrics.json").exists()


def test_cli_compare_same_dir_rejected(tmp_path):
    rc = main(["--config", str(FR4), "--results", str(tmp_path / "a"), "--compare", str(RO),
               "--compare-results", str(tmp_path / "a")])
    assert rc == 2
