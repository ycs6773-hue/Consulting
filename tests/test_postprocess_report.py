"""P3: CSV parsing, figures of merit, and report generation on synthetic PyAEDT-format data."""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import numpy as np
import pytest
from docx import Document

from hfss_patch import postprocess as pp
from hfss_patch.config import load_config
from hfss_patch.report import Spec, analyze, build_docx, main, verdicts
from hfss_patch.synthetic import write_synthetic_results

CFG = Path(__file__).resolve().parents[1] / "hfss_patch" / "configs" / "patch_2g4_fr4.yaml"


@pytest.fixture(scope="module")
def results(tmp_path_factory):
    cfg = load_config(CFG)
    return cfg, write_synthetic_results(cfg, tmp_path_factory.mktemp("syn"))


def test_header_parsing_units(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("Freq [MHz];L [cm];dB(S(P1,P1));S(P1,P1) (Real);S(P1,P1) (Imag) [ ]\n2400;2.94;-12;0.1;0.2\n")
    t = pp.read_table(p)
    assert t.col("Freq")[0] == pytest.approx(2.4)
    assert t.units["L"] == "mm" and t.col("L")[0] == pytest.approx(29.4)
    assert t.has("S(P1,P1) (Real)") and t.has("S(P1,P1) (Imag)")


def test_bad_csv_raises(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("Freq [GHz];dB(S(P1,P1))\n")
    with pytest.raises(pp.ResultsError):
        pp.read_table(p)
    with pytest.raises(pp.ResultsError):
        pp.read_table(tmp_path / "nope.csv")


def test_s11_metrics_analytic_notch():
    f = np.linspace(2.0, 2.8, 801)
    s = -30 + 25 * ((f - 2.44) / 0.04) ** 2
    s = np.minimum(s, 0)
    m = pp.s11_metrics(f, s)
    assert m.f_res_ghz == pytest.approx(2.44)
    # -30 + 25 x^2 = -10 -> x = sqrt(0.8) -> half-width 0.04*0.894
    assert m.bw10_mhz == pytest.approx(2 * 40 * np.sqrt(0.8), rel=0.01)
    assert m.ism_pass is False  # 71.6 MHz < 83.5 MHz ISM band
    wide = np.minimum(-30 + 25 * ((f - 2.442) / 0.06) ** 2, 0)
    assert pp.s11_metrics(f, wide).ism_pass is True


def test_s11_band_not_covered():
    f = np.linspace(2.0, 2.45, 100)
    assert pp.s11_metrics(f, np.full_like(f, -20)).ism_pass is False


def test_hpbw_cosine_pattern():
    a = np.arange(-180, 181, 1.0)
    g = 20 * np.log10(np.clip(np.cos(np.radians(a)), 1e-3, None))  # cos field -> HPBW 90 deg
    assert pp.hpbw(a, g) == pytest.approx(90.0, abs=0.5)


def test_tune_parameter():
    rows = [{"L": 28.0, "f_res_ghz": 2.5}, {"L": 30.0, "f_res_ghz": 2.3}]
    assert pp.tune_parameter(rows, "L", 2.4) == pytest.approx(29.0)


def test_synthetic_metrics(results):
    cfg, out = results
    m, curves = pp.s11_from_table(pp.read_table(out / "s11.csv"))
    assert 2.3 < m.f_res_ghz < 2.45 and m.s11_min_db < -20
    assert m.z_res_ohm is not None and abs(m.z_res_ohm.real - 50) < 5
    pm, cuts = pp.pattern_metrics(pp.read_table(out / "farfield_cuts.csv"))
    assert cuts["E"].angle_deg.min() < -170 and cuts["E"].angle_deg.max() == 180
    assert pm.hpbw_e_deg and pm.hpbw_h_deg and pm.front_to_back_db > 10
    assert pp.peak_from_3d(pp.read_table(out / "farfield_3d.csv"))[1] == 0.0  # broadside


def test_report_end_to_end(results, tmp_path):
    cfg, out = results
    rd = analyze(cfg, out)
    assert rd.synthetic and rd.missing == []
    assert set(rd.parametric) == {"L", "y0", "gap"}
    assert 28 < rd.tuning["L_for_f0_mm"] < 31
    v = {item: ok for item, _, _, ok in verdicts(rd, Spec())}
    assert v["ISM 대역 내 최악 |S11|"] is False  # FR-4 1.6 mm cannot cover 83.5 MHz
    doc_path = build_docx(rd, tmp_path / "r.docx")
    doc = Document(str(doc_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "합성(synthetic)" in text and "NOT FOR DELIVERY" in doc.sections[0].header.paragraphs[0].text
    assert len(doc.inline_shapes) == 8


def test_report_missing_files_degrade(tmp_path):
    cfg = load_config(CFG)
    out = write_synthetic_results(cfg, tmp_path / "r")
    (out / "farfield_3d.csv").unlink()
    (out / "s11_param_gap.csv").write_text("garbage\n")
    rd = analyze(cfg, out)
    assert any("farfield_3d" in m for m in rd.missing) and any("gap" in m for m in rd.missing)
    build_docx(rd, tmp_path / "r.docx")  # still builds


def test_real_results_not_watermarked(tmp_path):
    cfg = load_config(CFG)
    out = write_synthetic_results(cfg, tmp_path / "r")
    s = json.loads((out / "run_summary.json").read_text())
    s.pop("data_origin")
    (out / "run_summary.json").write_text(json.dumps(s))
    doc = Document(str(build_docx(analyze(cfg, out), tmp_path / "r.docx")))
    assert "SYNTHETIC" not in doc.sections[0].header.paragraphs[0].text


def test_sha_mismatch_flagged(tmp_path):
    cfg = load_config(CFG)
    out = write_synthetic_results(replace(cfg, sha256="0" * 64), tmp_path / "r")
    assert any("sha256" in m for m in analyze(cfg, out).missing)


def test_cli_synthetic(tmp_path):
    rc = main(["--config", str(CFG), "--results", str(tmp_path / "x"), "--synthetic"])
    assert rc == 0
    assert (tmp_path / "x" / "report.docx").exists() and (tmp_path / "x" / "metrics.json").exists()
